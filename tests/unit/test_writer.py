from __future__ import annotations

import datetime as dt
import errno
import os
import stat
import sys
from collections.abc import Sequence
from pathlib import Path
from typing import Any

import pytest

from cleandata.core.ledger import Reject
from cleandata.core.writer import (
    EXCEL_ROW_LIMIT,
    AtomicFile,
    CsvSink,
    ParquetSink,
    RejectsSink,
    check_output_path,
    parquet_available,
    reject_columns,
    save_workbook,
)
from cleandata.errors import MissingExtraError, OutputError, UsageError
from support import frame_of

COLUMNS = ["id", "when", "at", "n", "x", "ok", "note"]
ROWS = [
    ["a", dt.date(2025, 3, 21), dt.datetime(2025, 3, 22, 8, 1), 45000, 0.1, True, 'say "hi", twice'],
    ["b", None, None, None, None, False, "line\nbreak"],
]


def test_canonical_csv_bytes(tmp_path: Path) -> None:
    target = tmp_path / "out.csv"
    sink = CsvSink(target, COLUMNS)
    sink.write(frame_of(COLUMNS, ROWS))
    sink.commit()
    assert target.read_bytes() == (
        b"id,when,at,n,x,ok,note\n"
        b'a,2025-03-21,2025-03-22T08:01:00,45000,0.1,true,"say ""hi"", twice"\n'
        b'b,,,,,false,"line\nbreak"\n'
    )


def test_bom_and_delimiter_options(tmp_path: Path) -> None:
    target = tmp_path / "out.csv"
    sink = CsvSink(target, ["a", "b"], delimiter=";", bom=True)
    sink.write_rows([["1;2", "x"]])
    sink.commit()
    assert target.read_bytes() == b'\xef\xbb\xbfa;b\n"1;2";x\n'


def test_an_unencodable_value_is_an_output_error_and_leaves_no_file(tmp_path: Path) -> None:
    target = tmp_path / "out.csv"
    sink = CsvSink(target, ["a"], encoding="latin-1")
    with pytest.raises(OutputError, match="cannot be written as latin-1"):
        sink.write_rows([["€"]])
    assert list(tmp_path.iterdir()) == []


def test_atomic_file_replaces_only_on_commit(tmp_path: Path) -> None:
    target = tmp_path / "out.csv"
    target.write_text("old", encoding="utf-8")
    file = AtomicFile(target)
    file.temp.write_text("new", encoding="utf-8")
    assert target.read_text(encoding="utf-8") == "old"
    file.abort()
    assert target.read_text(encoding="utf-8") == "old"
    assert [p.name for p in tmp_path.iterdir()] == ["out.csv"]
    file = AtomicFile(target)
    file.temp.write_text("new", encoding="utf-8")
    file.commit()
    assert target.read_text(encoding="utf-8") == "new"


def test_output_path_checks(tmp_path: Path) -> None:
    with pytest.raises(UsageError, match="use one of .csv"):
        check_output_path(tmp_path / "out.txt", (".csv", ".xlsx"), "--out")
    folder = tmp_path / "folder.csv"
    folder.mkdir()
    with pytest.raises(OutputError, match="is a directory"):
        check_output_path(folder, (".csv",), "--out")


def test_xlsx_header_freeze_filter_and_formats(tmp_path: Path) -> None:
    import openpyxl

    target = tmp_path / "out.xlsx"
    save_workbook(target, "cleaned", COLUMNS, [*ROWS, ["=1+1", None, None, 1, 2.5, None, "x" * 80]])
    book = openpyxl.load_workbook(target)
    sheet = book["cleaned"]
    assert [cell.value for cell in sheet[1]] == COLUMNS
    assert all(cell.font.bold for cell in sheet[1])
    assert sheet.freeze_panes == "A2"
    assert sheet.auto_filter.ref == "A1:G4"
    assert sheet["B2"].value == dt.datetime(2025, 3, 21)
    assert sheet["B2"].number_format == "yyyy-mm-dd"
    assert sheet["C2"].number_format == "yyyy-mm-dd hh:mm:ss"
    assert sheet["D2"].value == 45000
    assert sheet["F2"].value is True
    assert sheet["A4"].value == "=1+1"
    assert sheet["A4"].data_type == "s"
    assert sheet.column_dimensions["G"].width == 50


def test_xlsx_control_characters_are_an_output_error(tmp_path: Path) -> None:
    with pytest.raises(OutputError, match="control character"):
        save_workbook(tmp_path / "out.xlsx", "cleaned", ["a"], [["bell\x07"]])


@pytest.mark.needs_parquet
def test_parquet_types(tmp_path: Path) -> None:
    import pyarrow.parquet as pq

    target = tmp_path / "out.parquet"
    types = {"id": "string", "when": "date", "at": "datetime", "n": "integer", "x": "number", "ok": "boolean"}
    sink = ParquetSink(target, COLUMNS, types)
    sink.write(frame_of(COLUMNS, ROWS))
    sink.commit()
    table = pq.read_table(target)
    assert [str(field.type) for field in table.schema] == [
        "string",
        "date32[day]",
        "timestamp[us]",
        "int64",
        "double",
        "bool",
        "string",
    ]
    assert table.column("n").to_pylist() == [45000, None]
    assert table.column("when").to_pylist() == [dt.date(2025, 3, 21), None]
    assert pq.ParquetFile(target).metadata.row_group(0).column(0).compression == "SNAPPY"


def test_rejects_columns_and_rows(tmp_path: Path) -> None:
    assert reject_columns(["id"], workbook=False) == ["_source_row", "_stage", "_rule", "_detail", "id"]
    assert reject_columns(["id"], workbook=True)[:2] == ["_source_row", "_source_sheet"]
    target = tmp_path / "rejects.csv"
    sink = RejectsSink(target, ["id", "age"], workbook=False)
    sink.write(
        [
            Reject(5, 7, None, "validate", ["age.max"], "age.max: 130 > 99", {"id": "NG-1", "age": "130"}),
            Reject(1, 3, None, "read", ["read.malformed"], "a;b;c;d"),
        ]
    )
    sink.commit()
    assert target.read_text(encoding="utf-8") == (
        "_source_row,_stage,_rule,_detail,id,age\n"
        "3,read,read.malformed,a;b;c;d,,\n"
        "7,validate,age.max,age.max: 130 > 99,NG-1,130\n"
    )


def test_xlsx_headers_and_error_like_text_stay_text(tmp_path: Path) -> None:
    import openpyxl

    target = tmp_path / "out.xlsx"
    save_workbook(target, "cleaned", ["=A1+1", "b"], [["#N/A", "#REF!"], ["=1+1", "ok"]])
    sheet = openpyxl.load_workbook(target)["cleaned"]
    cells = [(cell.value, cell.data_type) for line in sheet.iter_rows() for cell in line]
    assert cells == [
        ("=A1+1", "s"),
        ("b", "s"),
        ("#N/A", "s"),
        ("#REF!", "s"),
        ("=1+1", "s"),
        ("ok", "s"),
    ]


class HugeRows(Sequence[Sequence[object]]):
    """A row list of exactly the Excel limit that fails if anything reads a row."""

    def __len__(self) -> int:
        return EXCEL_ROW_LIMIT

    def __getitem__(self, index: Any) -> Any:
        raise AssertionError("the row limit is checked before any row is read")


def test_more_rows_than_excel_holds_is_an_output_error(tmp_path: Path) -> None:
    with pytest.raises(OutputError, match=f"exceed the Excel limit of {EXCEL_ROW_LIMIT - 1} data rows"):
        save_workbook(tmp_path / "out.xlsx", "cleaned", ["a"], HugeRows())
    assert list(tmp_path.iterdir()) == []


def test_xlsx_write_failure_is_an_output_error_and_leaves_no_file(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    def full_disk(self: Path, data: bytes) -> int:
        raise OSError(errno.ENOSPC, "No space left on device")

    with monkeypatch.context() as patch:
        patch.setattr(Path, "write_bytes", full_disk)
        with pytest.raises(OutputError, match="No space left on device"):
            save_workbook(tmp_path / "out.xlsx", "cleaned", ["a"], [["x"]])
    assert list(tmp_path.iterdir()) == []


def test_abort_ignores_a_temp_file_that_cannot_be_removed(tmp_path: Path) -> None:
    class Stuck:
        def unlink(self, missing_ok: bool = False) -> None:
            raise PermissionError(errno.EACCES, "Permission denied")

    file = AtomicFile(tmp_path / "out.csv")
    real = file.temp
    file.temp = Stuck()  # stands in for a Path whose unlink fails
    file.abort()
    real.unlink()
    assert list(tmp_path.iterdir()) == []


def test_parquet_output_without_pyarrow_is_a_missing_extra_error(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setitem(sys.modules, "pyarrow", None)
    assert not parquet_available()
    with pytest.raises(MissingExtraError, match='optional extra "parquet"'):
        ParquetSink(tmp_path / "out.parquet", ["a"], {})
    assert list(tmp_path.iterdir()) == []


@pytest.mark.needs_parquet
def test_a_parquet_writer_that_cannot_open_leaves_no_temp_file(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    import pyarrow.parquet as pq

    def broken(*args: object, **kwargs: object) -> None:
        raise OSError(errno.EACCES, "Permission denied")

    monkeypatch.setattr(pq, "ParquetWriter", broken)
    with pytest.raises(OutputError, match="Permission denied"):
        ParquetSink(tmp_path / "out.parquet", ["a"], {})
    assert list(tmp_path.iterdir()) == []


class FullDiskWriter:
    def writerow(self, row: object) -> None:
        raise OSError(errno.ENOSPC, "No space left on device")


class FullDiskHandle:
    def __init__(self, real: Any) -> None:
        self.real = real

    def flush(self) -> None:
        raise OSError(errno.ENOSPC, "No space left on device")

    def close(self) -> None:
        self.real.close()


def test_a_failing_csv_write_is_an_output_error_and_keeps_the_old_file(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    target = tmp_path / "out.csv"
    target.write_text("old", encoding="utf-8")
    sink = CsvSink(target, ["a"])
    monkeypatch.setattr(sink, "writer", FullDiskWriter())
    with pytest.raises(OutputError, match="No space left on device"):
        sink.write_rows([["x"]])
    assert target.read_text(encoding="utf-8") == "old"
    assert [p.name for p in tmp_path.iterdir()] == ["out.csv"]


def test_a_failing_csv_flush_on_commit_is_an_output_error_and_leaves_no_file(tmp_path: Path) -> None:
    sink = CsvSink(tmp_path / "out.csv", ["a"])
    sink.handle = FullDiskHandle(sink.handle)
    with pytest.raises(OutputError, match="No space left on device"):
        sink.commit()
    assert list(tmp_path.iterdir()) == []


def test_a_csv_sink_that_cannot_start_leaves_no_temp_file(tmp_path: Path) -> None:
    with pytest.raises(LookupError):
        CsvSink(tmp_path / "out.csv", ["a"], encoding="no-such-codec")
    assert list(tmp_path.iterdir()) == []


def spy_on_chmod(monkeypatch: pytest.MonkeyPatch) -> list[int]:
    modes: list[int] = []
    real = os.chmod

    def spy(path: Any, mode: int, **kwargs: Any) -> None:
        modes.append(mode)
        real(path, mode, **kwargs)

    monkeypatch.setattr(os, "chmod", spy)
    return modes


def test_a_new_file_gets_the_mode_a_plain_open_would_give(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    modes = spy_on_chmod(monkeypatch)
    target = tmp_path / "out.csv"
    previous = os.umask(0o027)
    try:
        kept = os.umask(0o027)  # Windows keeps only the owner bits of the mask
        sink = CsvSink(target, ["a"])
        sink.commit()
    finally:
        os.umask(previous)
    assert modes == [0o666 & ~kept]
    if os.name == "posix":
        assert stat.S_IMODE(target.stat().st_mode) == 0o640
    assert os.access(target, os.R_OK)
    assert target.stat().st_mode & stat.S_IWRITE


def test_a_replaced_file_keeps_the_mode_of_the_old_one(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    target = tmp_path / "out.csv"
    target.write_text("old", encoding="utf-8")
    target.chmod(0o640)  # on Windows this only keeps the file writable
    expected = stat.S_IMODE(target.stat().st_mode)
    modes = spy_on_chmod(monkeypatch)
    sink = CsvSink(target, ["a"])
    sink.commit()
    assert modes == [expected]
    assert stat.S_IMODE(target.stat().st_mode) == expected
    assert os.access(target, os.R_OK)
    assert target.stat().st_mode & stat.S_IWRITE
