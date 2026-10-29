"""Stage 1, read: decode, split records, find the header row, skip footer rows, number the rows.

The CSV reader sits on the stdlib ``csv`` module and keeps the physical line of every record
(``_source_row``: 1-based, header counted, so it matches a text editor for single-line records and
Excel's row number for workbooks). A record shorter than the header is padded with nulls; a longer
one is accepted when its extra fields are empty and rejected as ``read.malformed`` otherwise, with
its raw text in ``_detail``. Workbooks are read with openpyxl in read-only mode with cached formula
values; cells keep their native int, float, bool and datetime values.
"""

from __future__ import annotations

import csv
import fnmatch
import hashlib
import zipfile
from collections import deque
from collections.abc import Generator, Iterator, Sequence
from dataclasses import dataclass, field
from pathlib import Path
from typing import TYPE_CHECKING, Any, Final

import openpyxl
import pandas as pd

from cleandata.core.coerce import canonical_str
from cleandata.core.frame import META_ROW, META_SHEET, new_frame, positional
from cleandata.core.ledger import Reject
from cleandata.core.sniff import (
    NO_DELIMITER,
    detect_header_row,
    find_decode_error,
    is_blank,
    resolve_encoding,
    sample_text,
    sniff_delimiter,
)
from cleandata.errors import InputError

if TYPE_CHECKING:
    from cleandata.core.profile import Profile

WORKBOOK_SUFFIXES: Final = (".xlsx", ".xlsm")
UNSUPPORTED_SUFFIXES: Final = (".xls", ".ods", ".json", ".xml", ".parquet")
RAW_DETAIL_LIMIT: Final = 500
FIELD_LIMIT: Final = 2**31 - 1


@dataclass(frozen=True)
class ReadOptions:
    encoding: str = "auto"
    delimiter: str = "auto"
    quotechar: str = '"'
    header_row: int | str = 1
    skip_footer_rows: int = 0
    sheets: tuple[str, ...] | None = None

    @classmethod
    def from_profile(cls, profile: Profile) -> ReadOptions:
        options = profile.input
        return cls(
            encoding=options.encoding,
            delimiter=options.delimiter,
            quotechar=options.quotechar,
            header_row=options.header_row,
            skip_footer_rows=options.skip_footer_rows,
            sheets=tuple(options.sheets) if options.sheets else None,
        )


@dataclass
class SheetInfo:
    name: str
    header_row: int
    records: int


@dataclass
class SourceInfo:
    name: str
    size: int
    sha256: str
    kind: str
    encoding: str | None = None
    encoding_method: str | None = None
    encoding_confidence: float | None = None
    delimiter: str | None = None
    delimiter_share: float | None = None
    header_row: int | None = None
    sheets: list[SheetInfo] = field(default_factory=list)
    ignored_sheets: list[str] = field(default_factory=list)
    records: int = 0
    footer_skipped: int = 0
    short_records: int = 0
    malformed: int = 0


@dataclass
class RawChunk:
    """Raw cells in positional columns ``#0`` ... plus the meta columns; index = record ordinal."""

    frame: pd.DataFrame
    rejects: list[Reject]
    read: int


def file_digest(path: Path) -> tuple[int, str]:
    digest = hashlib.sha256()
    size = 0
    with path.open("rb") as handle:
        while block := handle.read(1 << 20):
            digest.update(block)
            size += len(block)
    return size, digest.hexdigest()


class _LineCounter:
    """Feeds lines to ``csv.reader`` and remembers which lines made up the current record."""

    def __init__(self, lines: Iterator[str]) -> None:
        self._lines = lines
        self.line = 0
        self.buffer: list[str] = []

    def __iter__(self) -> _LineCounter:
        return self

    def __next__(self) -> str:
        text = next(self._lines)
        self.line += 1
        self.buffer.append(text)
        return text


class CsvSource:
    def __init__(self, path: Path, options: ReadOptions, info: SourceInfo) -> None:
        self.path = path
        self.options = options
        self.info = info
        guess = resolve_encoding(path, options.encoding)
        self.encoding = guess.encoding
        info.encoding, info.encoding_method, info.encoding_confidence = (
            guess.encoding,
            guess.method,
            guess.confidence,
        )
        if options.delimiter != "auto":
            self.delimiter = options.delimiter
            info.delimiter, info.delimiter_share = options.delimiter, None
        else:
            sniffed = sniff_delimiter(sample_text(path, self.encoding), options.quotechar)
            self.delimiter = sniffed.delimiter or NO_DELIMITER
            info.delimiter, info.delimiter_share = sniffed.delimiter, sniffed.share
        wanted = options.header_row if isinstance(options.header_row, int) else 20
        head: list[list[str]] = []
        records = self._records()
        try:
            for _, fields, _ in records:
                head.append(fields)
                if len(head) >= wanted:
                    break
        finally:
            records.close()
        if isinstance(options.header_row, int):
            number: int | None = options.header_row if len(head) >= options.header_row else None
        else:
            number = detect_header_row(head)
        if number is None:
            raise InputError(f"{path.name}: header row {options.header_row} not found")
        self.header_row = number
        info.header_row = number
        self.headers = [name.strip() for name in head[number - 1]]

    def _records(self) -> Generator[tuple[int, list[str], str], None, None]:
        csv.field_size_limit(FIELD_LIMIT)
        try:
            with self.path.open("r", encoding=self.encoding, errors="strict", newline="") as handle:
                counter = _LineCounter(iter(handle))
                reader = csv.reader(
                    counter, delimiter=self.delimiter, quotechar=self.options.quotechar, strict=False
                )
                for fields in reader:
                    start = counter.line - len(counter.buffer) + 1
                    raw = "".join(counter.buffer)
                    counter.buffer.clear()
                    yield start, fields, raw
        except UnicodeDecodeError as exc:
            offset = find_decode_error(self.path, self.encoding)
            raise InputError(
                f"{self.path.name}: cannot decode as {self.encoding} at byte offset {offset}; "
                "set input.encoding in the profile"
            ) from exc
        except csv.Error as exc:
            raise InputError(f"{self.path.name}: {exc}") from exc

    def chunks(self, size: int | None) -> Iterator[RawChunk]:
        width = len(self.headers)
        names = [positional(i) for i in range(width)] + [META_ROW, META_SHEET]
        footer = self.options.skip_footer_rows
        pending: deque[tuple[int, list[str], str]] = deque()
        rows: list[list[object]] = []
        index: list[int] = []
        rejects: list[Reject] = []
        read = 0
        ordinal = 0
        for number, record in enumerate(self._records(), start=1):
            if number <= self.header_row:
                continue
            pending.append(record)
            if len(pending) <= footer:
                continue
            line, fields, raw = pending.popleft()
            read += 1
            values: list[object]
            if all(value.strip() == "" for value in fields):
                values = [""] * width
            elif len(fields) < width:
                values = [*fields, *([None] * (width - len(fields)))]
                self.info.short_records += 1
            elif len(fields) > width and any(value.strip() for value in fields[width:]):
                detail = raw.rstrip("\r\n")[:RAW_DETAIL_LIMIT]
                rejects.append(Reject(ordinal, line, None, "read", ["read.malformed"], detail))
                self.info.malformed += 1
                ordinal += 1
                values = []
            else:
                values = list(fields[:width])
            if values:
                rows.append([*values, line, None])
                index.append(ordinal)
                ordinal += 1
            if size is not None and read >= size:
                yield RawChunk(new_frame(names, rows, index), rejects, read)
                self.info.records += read
                rows, index, rejects, read = [], [], [], 0
        self.info.footer_skipped = len(pending)
        self.info.records += read
        yield RawChunk(new_frame(names, rows, index), rejects, read)


@dataclass
class _Sheet:
    name: str
    header_row: int
    width: int
    union: list[int]
    records: list[tuple[int, list[object]]]


def _cells(row: Sequence[object]) -> list[object]:
    return list(row)


class WorkbookSource:
    def __init__(self, path: Path, options: ReadOptions, info: SourceInfo) -> None:
        self.path = path
        self.options = options
        self.info = info
        try:
            book: Any = openpyxl.load_workbook(path, read_only=True, data_only=True)
        except (zipfile.BadZipFile, KeyError, OSError, ValueError) as exc:
            raise InputError(f"{path.name}: not a readable workbook ({exc})") from exc
        try:
            names = [str(name) for name in book.sheetnames]
            selected = self._select(names)
            sheet_rows: dict[str, list[list[object]]] = {}
            for name in selected:
                worksheet = book[name]
                worksheet.reset_dimensions()
                sheet_rows[name] = [_cells(row) for row in worksheet.iter_rows(values_only=True)]
        finally:
            book.close()
        info.ignored_sheets = [name for name in names if name not in selected]
        keys: list[tuple[str, int]] = []
        self.sheets: list[_Sheet] = []
        for name in selected:
            rows = sheet_rows[name]
            while rows and all(is_blank(value) for value in rows[-1]):
                rows.pop()
            if isinstance(options.header_row, int):
                number: int | None = options.header_row if len(rows) >= options.header_row else None
            else:
                number = detect_header_row(rows)
            if number is None:
                raise InputError(f"{path.name}: header row {options.header_row} not found on sheet {name!r}")
            header = ["" if value is None else canonical_str(value).strip() for value in rows[number - 1]]
            while header and header[-1] == "":
                header.pop()
            seen: dict[str, int] = {}
            union: list[int] = []
            for text in header:
                seen[text] = seen.get(text, 0) + 1
                key = (text, seen[text])
                if key not in keys:
                    keys.append(key)
                union.append(keys.index(key))
            body = list(enumerate(rows, start=1))[number:]
            footer = options.skip_footer_rows
            if footer:
                info.footer_skipped += min(footer, len(body))
                body = body[: max(0, len(body) - footer)]
            self.sheets.append(_Sheet(name, number, len(header), union, body))
            info.sheets.append(SheetInfo(name, number, len(body)))
        self.headers = [text for text, _ in keys]

    def _select(self, names: list[str]) -> list[str]:
        patterns = self.options.sheets
        if not patterns:
            return names[:1]
        selected = [name for name in names if any(fnmatch.fnmatchcase(name, pattern) for pattern in patterns)]
        if not selected:
            raise InputError(
                f"{self.path.name}: no sheet matches {list(patterns)} (sheets: {', '.join(names)})"
            )
        return selected

    def chunks(self, size: int | None) -> Iterator[RawChunk]:
        width = len(self.headers)
        names = [positional(i) for i in range(width)] + [META_ROW, META_SHEET]
        rows: list[list[object]] = []
        index: list[int] = []
        rejects: list[Reject] = []
        ordinal = 0
        read = 0
        for sheet in self.sheets:
            for number, cells in sheet.records:
                read += 1
                extra = cells[sheet.width :]
                if any(not is_blank(value) for value in extra):
                    detail = "\t".join(canonical_str(value) for value in cells)[:RAW_DETAIL_LIMIT]
                    rejects.append(Reject(ordinal, number, sheet.name, "read", ["read.malformed"], detail))
                    self.info.malformed += 1
                    ordinal += 1
                    continue
                values: list[object] = [None] * width
                for local, value in enumerate(cells[: sheet.width]):
                    values[sheet.union[local]] = value
                rows.append([*values, number, sheet.name])
                index.append(ordinal)
                ordinal += 1
        self.info.records = read
        yield RawChunk(new_frame(names, rows, index), rejects, read)


Source = CsvSource | WorkbookSource


def open_source(path: Path, options: ReadOptions) -> Source:
    """Open a CSV-family file or a workbook; raises ``InputError`` (exit 3) when it cannot be read."""
    if not path.is_file():
        raise InputError(f"{path}: no such file")
    suffix = path.suffix.lower()
    if suffix in UNSUPPORTED_SUFFIXES:
        raise InputError(f"{path.name}: {suffix} files are not supported (CSV, TXT, TSV, XLSX, XLSM)")
    try:
        size, digest = file_digest(path)
    except OSError as exc:
        raise InputError(f"{path.name}: cannot read the file ({exc.strerror or exc})") from exc
    kind = "xlsx" if suffix in WORKBOOK_SUFFIXES else "csv"
    info = SourceInfo(name=path.name, size=size, sha256=digest, kind=kind)
    if kind == "xlsx":
        return WorkbookSource(path, options, info)
    return CsvSource(path, options, info)
