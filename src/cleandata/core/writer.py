"""Stage 8, write: canonical CSV, XLSX or (extra) Parquet, the rejects file and atomic replacement.

The canonical CSV is UTF-8 without BOM unless ``output.csv.bom``, comma, ``"`` quoting only where
needed, ``\\n`` line ends on every OS, a header row, null as an empty field, ISO dates and datetimes,
integers without decimals, numbers as Python's shortest round-trip repr and booleans ``true`` and
``false``. It is written with the stdlib ``csv`` module, not ``DataFrame.to_csv``, so both drivers
produce the same bytes. Every file goes to a temporary file in the target folder and is moved into
place with ``os.replace``, so a failed run never leaves half a file.
"""

from __future__ import annotations

import codecs
import csv
import datetime as dt
import importlib
import importlib.util
import os
import tempfile
from collections.abc import Mapping, Sequence
from pathlib import Path
from typing import IO, Any, Final, Protocol

import openpyxl
import pandas as pd
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from openpyxl.utils.exceptions import IllegalCharacterError

from cleandata.core.coerce import canonical_str
from cleandata.core.frame import rows_of
from cleandata.core.ledger import Reject
from cleandata.errors import MissingExtraError, OutputError, UsageError

EXCEL_ROW_LIMIT: Final = 1_048_576
OUTPUT_SUFFIXES: Final = (".csv", ".xlsx", ".parquet")
REJECTS_SUFFIXES: Final = (".csv", ".xlsx")
REPORT_SUFFIXES: Final = (".html", ".json")
META_REJECT_COLUMNS: Final = ("_stage", "_rule", "_detail")
MAX_WIDTH: Final = 50


def parquet_available() -> bool:
    return importlib.util.find_spec("pyarrow") is not None


class AtomicFile:
    """A temporary file next to ``target``; ``commit`` moves it into place, ``abort`` removes it."""

    def __init__(self, target: Path) -> None:
        self.target = target
        try:
            target.parent.mkdir(parents=True, exist_ok=True)
            handle, name = tempfile.mkstemp(prefix=f".{target.name}.", suffix=".tmp", dir=target.parent)
        except OSError as exc:
            raise OutputError(f"cannot write {target}: {exc.strerror or exc}") from exc
        os.close(handle)
        self.temp = Path(name)

    def commit(self) -> None:
        try:
            os.replace(self.temp, self.target)
        except OSError as exc:
            self.abort()
            raise OutputError(f"cannot write {self.target}: {exc.strerror or exc}") from exc

    def abort(self) -> None:
        self.temp.unlink(missing_ok=True)


def write_text(target: Path, text: str) -> None:
    """Write UTF-8 text with ``\\n`` line ends atomically."""
    file = AtomicFile(target)
    try:
        with file.temp.open("w", encoding="utf-8", newline="\n") as handle:
            handle.write(text)
    except BaseException:
        file.abort()
        raise
    file.commit()


class Sink(Protocol):
    def write(self, frame: pd.DataFrame) -> None: ...

    def commit(self) -> None: ...

    def abort(self) -> None: ...


def _csv_encoding(encoding: str, bom: bool) -> str:
    return "utf-8-sig" if bom and codecs.lookup(encoding).name == "utf-8" else encoding


class CsvSink:
    def __init__(
        self,
        target: Path,
        columns: Sequence[str],
        delimiter: str = ",",
        encoding: str = "utf-8",
        bom: bool = False,
    ) -> None:
        self.file = AtomicFile(target)
        self.columns = list(columns)
        self.encoding = _csv_encoding(encoding, bom)
        self.handle: IO[str] = self.file.temp.open("w", encoding=self.encoding, newline="")
        self.writer = csv.writer(self.handle, delimiter=delimiter, quotechar='"', lineterminator="\n")
        self._row(self.columns)

    def _row(self, values: Sequence[str]) -> None:
        try:
            self.writer.writerow(values)
        except UnicodeEncodeError as exc:
            self.abort()
            raise OutputError(
                f"{self.file.target.name}: a value cannot be written as {self.encoding}"
            ) from exc

    def write(self, frame: pd.DataFrame) -> None:
        for row in rows_of(frame, self.columns):
            self._row([canonical_str(value) for value in row])

    def write_rows(self, rows: Sequence[Sequence[str]]) -> None:
        for row in rows:
            self._row(row)

    def commit(self) -> None:
        self.handle.close()
        self.file.commit()

    def abort(self) -> None:
        self.handle.close()
        self.file.abort()


def _xlsx_cell(value: object) -> object:
    if value is None or isinstance(value, str | bool | int | float | dt.date | dt.datetime):
        return value
    return canonical_str(value)


def save_workbook(
    target: Path, sheet_name: str, columns: Sequence[str], rows: Sequence[Sequence[object]]
) -> None:
    """One sheet with a bold frozen header, an autofilter, widths from content and ISO date formats."""
    if len(rows) + 1 > EXCEL_ROW_LIMIT:
        raise OutputError(
            f"{target.name}: {len(rows)} rows exceed the Excel limit of {EXCEL_ROW_LIMIT - 1} data rows; "
            "write CSV or Parquet instead"
        )
    book: Any = openpyxl.Workbook()
    sheet = book.active
    sheet.title = sheet_name
    widths = [len(name) for name in columns]
    try:
        sheet.append(list(columns))
        for row in rows:
            cells = [_xlsx_cell(value) for value in row]
            sheet.append(cells)
            for index, value in enumerate(cells):
                widths[index] = max(widths[index], len(canonical_str(value)))
    except IllegalCharacterError as exc:
        raise OutputError(
            f"{target.name}: a cell holds a control character that XLSX cannot store; "
            "keep text.remove_control on"
        ) from exc
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for line in sheet.iter_rows(min_row=2):
        for cell in line:
            value = cell.value
            if isinstance(value, dt.datetime):
                cell.number_format = "yyyy-mm-dd hh:mm:ss"
            elif isinstance(value, dt.date):
                cell.number_format = "yyyy-mm-dd"
            elif isinstance(value, str) and value.startswith("="):
                cell.data_type = "s"
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = min(MAX_WIDTH, width + 2)
    sheet.freeze_panes = "A2"
    last = get_column_letter(max(1, len(columns)))
    sheet.auto_filter.ref = f"A1:{last}{len(rows) + 1}"
    file = AtomicFile(target)
    try:
        book.save(file.temp)
    except BaseException:
        file.abort()
        raise
    file.commit()


class XlsxSink:
    """Collects the rows (workbook output is in-memory only) and saves them on commit."""

    def __init__(self, target: Path, columns: Sequence[str], sheet_name: str) -> None:
        self.target = target
        self.columns = list(columns)
        self.sheet_name = sheet_name
        self.rows: list[tuple[object, ...]] = []

    def write(self, frame: pd.DataFrame) -> None:
        self.rows.extend(rows_of(frame, self.columns))

    def commit(self) -> None:
        save_workbook(self.target, self.sheet_name, self.columns, self.rows)

    def abort(self) -> None:
        self.rows = []


PARQUET_TYPES: Final = {
    "string": "string",
    "integer": "int64",
    "number": "float64",
    "boolean": "bool_",
    "date": "date32",
    "datetime": "timestamp",
}


class ParquetSink:
    """pyarrow's ParquetWriter: one row group per chunk, schema fixed from the profile types."""

    def __init__(self, target: Path, columns: Sequence[str], types: Mapping[str, str]) -> None:
        if not parquet_available():
            raise MissingExtraError("parquet", "Parquet output")
        self.pa: Any = importlib.import_module("pyarrow")
        parquet: Any = importlib.import_module("pyarrow.parquet")
        self.columns = list(columns)
        fields = []
        for name in self.columns:
            kind = PARQUET_TYPES[types.get(name, "string")]
            arrow_type = self.pa.timestamp("us") if kind == "timestamp" else getattr(self.pa, kind)()
            fields.append(self.pa.field(name, arrow_type))
        self.schema = self.pa.schema(fields)
        self.file = AtomicFile(target)
        self.writer = parquet.ParquetWriter(str(self.file.temp), self.schema, compression="snappy")

    def write(self, frame: pd.DataFrame) -> None:
        values = list(zip(*rows_of(frame, self.columns), strict=True)) if self.columns else []
        arrays = [
            self.pa.array(list(values[index]) if values else [], type=self.schema.field(index).type)
            for index in range(len(self.columns))
        ]
        self.writer.write_table(self.pa.Table.from_arrays(arrays, schema=self.schema))

    def commit(self) -> None:
        self.writer.close()
        self.file.commit()

    def abort(self) -> None:
        self.writer.close()
        self.file.abort()


def check_output_path(path: Path, allowed: Sequence[str], what: str) -> None:
    if path.suffix.lower() not in allowed:
        raise UsageError(f"{what} {path.name}: use one of {', '.join(allowed)}")
    if path.is_dir():
        raise OutputError(f"{what} {path} is a directory")


def open_output(
    path: Path,
    columns: Sequence[str],
    types: Mapping[str, str],
    csv_delimiter: str,
    csv_encoding: str,
    csv_bom: bool,
    sheet_name: str,
) -> Sink:
    suffix = path.suffix.lower()
    if suffix == ".csv":
        return CsvSink(path, columns, csv_delimiter, csv_encoding, csv_bom)
    if suffix == ".xlsx":
        return XlsxSink(path, columns, sheet_name)
    return ParquetSink(path, columns, types)


def reject_columns(data: Sequence[str], workbook: bool) -> list[str]:
    """``_source_row``, ``_source_sheet`` (workbook input only), ``_stage``, ``_rule``, ``_detail``, data."""
    return ["_source_row", *(["_source_sheet"] if workbook else []), *META_REJECT_COLUMNS, *data]


def reject_row(reject: Reject, data: Sequence[str], workbook: bool) -> list[str]:
    meta = [str(reject.source_row), *([reject.source_sheet or ""] if workbook else [])]
    return [
        *meta,
        reject.stage,
        reject.rule,
        reject.detail,
        *(canonical_str(reject.raw.get(name)) for name in data),
    ]


class RejectsSink:
    def __init__(self, path: Path, data: Sequence[str], workbook: bool) -> None:
        self.path = path
        self.data = list(data)
        self.workbook = workbook
        self.columns = reject_columns(self.data, workbook)
        self.csv = CsvSink(path, self.columns) if path.suffix.lower() == ".csv" else None
        self.rows: list[list[str]] = []

    def write(self, rejects: Sequence[Reject]) -> None:
        rows = [
            reject_row(item, self.data, self.workbook) for item in sorted(rejects, key=lambda r: r.ordinal)
        ]
        if self.csv is not None:
            self.csv.write_rows(rows)
        else:
            self.rows.extend(rows)

    def commit(self) -> None:
        if self.csv is not None:
            self.csv.commit()
        else:
            save_workbook(self.path, "rejects", self.columns, self.rows)

    def abort(self) -> None:
        if self.csv is not None:
            self.csv.abort()
