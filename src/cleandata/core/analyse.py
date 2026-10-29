"""The ``profile`` command: what is in a file before any profile exists, and the guesses ``init`` uses.

Per column: header, counts, null-like tokens, distinct and top values, lengths, whitespace, non-NFC
and control characters, candidate types with their parse shares, and ambiguity notes. Per file:
duplicate headers, empty columns, raw duplicate records and a locale vote.
"""

from __future__ import annotations

import datetime as dt
import glob
import hashlib
import re
from collections import Counter
from collections.abc import Sequence
from dataclasses import dataclass, field
from functools import lru_cache
from pathlib import Path
from typing import Any, Final

import openpyxl

from cleandata.core.coerce import Invalid, ParseSettings, canonical_str, to_boolean, to_integer, to_number
from cleandata.core.frame import column, data_columns, row_count, rows_of
from cleandata.core.headers import header_bases, normalise_headers
from cleandata.core.reader import ReadOptions, SourceInfo, open_source
from cleandata.core.sniff import (
    BUILTIN_DATE_PATTERNS,
    BUILTIN_DATETIME_PATTERNS,
    NUMBER_STYLES,
    detect_header_row,
    is_blank,
)
from cleandata.core.stats import DISTINCT_CAP, TOP
from cleandata.core.text import TextSettings, clean_text, display, has_control, has_whitespace_issue, is_nfc
from cleandata.errors import InputError

DEFAULT_MAX_ROWS: Final = 200_000
GUESS_SHARE: Final = 0.95
NULL_LIKE: Final = ("-", "--", "?", "n/a", "na", "n.a.", "k.a.", "k. a.", "null", "none", "nan", "#n/a")
CURRENCY: Final = ("€", "£", "$")
DEFAULT_TEXT: Final = TextSettings()
EXAMPLES: Final = 3


@dataclass
class Candidate:
    type: str
    share: float
    decimal: str | None = None
    thousands: str | None = None
    patterns: list[str] = field(default_factory=list)
    strip: list[str] = field(default_factory=list)

    def label(self) -> str:
        if self.type in ("integer", "number"):
            grouping = "none" if self.thousands is None else repr(self.thousands)
            return f"{self.type}(decimal={self.decimal!r}, thousands={grouping})"
        if self.type in ("date", "datetime"):
            if not self.patterns and self.share > 0:
                return f"{self.type} (native)"  # every hit was a native Excel date, no text pattern
            return f"{self.type}({', '.join(self.patterns)})"
        return self.type

    def as_dict(self) -> dict[str, Any]:
        return {
            "type": self.type,
            "label": self.label(),
            "share": round(self.share, 4),
            "decimal": self.decimal,
            "thousands": self.thousands,
            "patterns": self.patterns,
            "strip": self.strip,
        }


@dataclass
class ColumnAnalysis:
    position: int
    raw: str
    normalised: str
    non_empty: int = 0
    null_like: dict[str, int] = field(default_factory=dict)
    distinct: int = 0
    distinct_capped: bool = False
    top: list[list[object]] = field(default_factory=list)
    min_length: int | None = None
    max_length: int | None = None
    whitespace: int = 0
    non_nfc: int = 0
    control: int = 0
    candidates: list[Candidate] = field(default_factory=list)
    guess: Candidate | None = None
    styles: dict[tuple[str, str | None], tuple[float, float]] = field(default_factory=dict)
    examples: list[str] = field(default_factory=list)
    notes: list[str] = field(default_factory=list)

    def as_dict(self) -> dict[str, Any]:
        return {
            "position": self.position,
            "raw": self.raw,
            "normalised": self.normalised,
            "non_empty": self.non_empty,
            "null_like": self.null_like,
            "distinct": self.distinct,
            "distinct_capped": self.distinct_capped,
            "top": self.top,
            "min_length": self.min_length,
            "max_length": self.max_length,
            "whitespace": self.whitespace,
            "non_nfc": self.non_nfc,
            "control": self.control,
            "candidates": [candidate.as_dict() for candidate in self.candidates],
            "guess": None if self.guess is None else self.guess.as_dict(),
            "examples": self.examples,
            "notes": self.notes,
        }


@dataclass
class Analysis:
    info: SourceInfo
    headers: list[str]
    columns: list[ColumnAnalysis]
    analysed: int
    empty: int
    stopped_early: bool
    duplicate_records: int
    duplicate_headers: list[str]
    empty_columns: list[str]
    locale: tuple[str, str | None]
    locale_votes: dict[str, int]
    date_formats: list[str]
    datetime_formats: list[str]
    all_sheets: list[tuple[str, int | None, int]]
    selected_sheets: list[str]
    warnings: list[str]


def _share(hits: int, total: int) -> float:
    return hits / total if total else 0.0


def _style_key(decimal: str, thousands: str | None) -> str:
    return f"decimal {decimal!r}, thousands " + ("none" if thousands is None else repr(thousands))


Weighted = Sequence[tuple[object, int]]


def weighted(values: Sequence[object]) -> list[tuple[object, int]]:
    """Distinct values with their counts, so every parse runs once per distinct value.

    The key carries the type, so ``True``, ``1`` and ``1.0`` from a workbook stay apart.
    """
    counts: Counter[tuple[str, object]] = Counter((type(value).__name__, value) for value in values)
    return [(value, count) for (_, value), count in counts.items()]


def _number_candidates(
    values: Weighted, total: int
) -> tuple[list[Candidate], list[str], dict[tuple[str, str | None], tuple[float, float]]]:
    """Number and integer shares per style, with currency symbols stripped when that helps."""
    notes: list[str] = []
    seen = [symbol for symbol in CURRENCY if any(isinstance(v, str) and symbol in v for v, _ in values)]
    strip = tuple(seen)
    results: list[Candidate] = []
    best_integer: Candidate | None = None
    styles: dict[tuple[str, str | None], tuple[float, float]] = {}
    for decimal, thousands in NUMBER_STYLES:
        settings = ParseSettings(decimal=decimal, thousands=thousands)
        numbers = integers = 0
        for value, count in values:
            if not isinstance(to_number(value, settings, strip), Invalid):
                numbers += count
                if not isinstance(to_integer(value, settings, strip), Invalid):
                    integers += count
        results.append(Candidate("number", _share(numbers, total), decimal, thousands, strip=list(seen)))
        candidate = Candidate("integer", _share(integers, total), decimal, thousands, strip=list(seen))
        styles[(decimal, thousands)] = (results[-1].share, candidate.share)
        if best_integer is None or candidate.share > best_integer.share:
            best_integer = candidate
    readings: dict[str, set[float]] = {}
    for value, _ in values:
        if not isinstance(value, str):
            continue
        parsed = {
            float(number)
            for decimal, thousands in ((".", None), (",", "."))
            if not isinstance(
                number := to_number(value, ParseSettings(decimal=decimal, thousands=thousands)), Invalid
            )
        }
        if len(parsed) > 1:
            readings.setdefault(value.strip(), parsed)
    if readings:
        example = next(iter(readings))
        notes.append(
            f'{len(readings)} values read differently in English and German style (e.g. "{example}" is '
            + " or ".join(f"{number:g}" for number in sorted(readings[example]))
            + ")"
        )
    best_number = max(results, key=lambda c: c.share)
    chosen = [best_number] + ([best_integer] if best_integer is not None else [])
    return chosen, notes, styles


DIRECTIVE_SHAPES: Final = {
    "%Y": r"\d{4}",
    "%y": r"\d{2}",
    "%m": r"\d{1,2}",
    "%d": r"\d{1,2}",
    "%H": r"\d{1,2}",
    "%M": r"\d{1,2}",
    "%S": r"\d{1,2}",
}


@lru_cache(maxsize=64)
def shape(pattern: str) -> re.Pattern[str]:
    """A cheap regular expression a value must match before ``strptime`` is tried with ``pattern``."""
    parts = re.split(r"(%[a-zA-Z])", pattern)
    return re.compile("".join(DIRECTIVE_SHAPES.get(part, re.escape(part)) for part in parts))


def _date_candidate(
    values: Weighted, total: int, patterns: Sequence[str], kind: str
) -> tuple[Candidate, list[str]]:
    counts: Counter[str] = Counter()
    hits = 0
    day_first = month_first = both = 0
    for value, count in values:
        if isinstance(value, dt.datetime):
            if kind == "datetime" or value.time() == dt.time(0, 0):
                hits += count
                counts["native"] += count
            continue
        if isinstance(value, dt.date):
            hits += count
            counts["native"] += count
            continue
        if not isinstance(value, str):
            continue
        text = value.strip()
        matched: list[str] = []
        for pattern in patterns:
            if shape(pattern).fullmatch(text) is None:
                continue
            try:
                dt.datetime.strptime(text, pattern)
            except ValueError:
                continue
            matched.append(pattern)
        if matched:
            hits += count
            counts[matched[0]] += count
            days = any(p.startswith("%d/%m/%Y") for p in matched)
            months = any(p.startswith("%m/%d/%Y") for p in matched)
            if days and months:
                both += count
            elif days:
                day_first += count
            elif months:
                month_first += count
    used = [pattern for pattern in patterns if counts[pattern]]
    notes: list[str] = []
    if both:
        notes.append(
            f"{both} values parse day-first and month-first (%d/%m/%Y and %m/%d/%Y); "
            + (
                "month-first is the only reading of some others"
                if month_first
                else "the guess reads them day-first"
            )
        )
    if month_first and day_first:
        notes.append("both %d/%m/%Y and %m/%d/%Y occur; check the source before trusting either")
    if month_first and not day_first:
        used = list(dict.fromkeys(p.replace("%d/%m/%Y", "%m/%d/%Y") for p in used))
    return Candidate(kind, _share(hits, total), patterns=used), notes


def _boolean_candidate(values: Weighted, total: int) -> Candidate:
    settings = ParseSettings()
    hits = sum(count for value, count in values if not isinstance(to_boolean(value, settings), Invalid))
    return Candidate("boolean", _share(hits, total))


def choose_guess(candidates: Sequence[Candidate]) -> Candidate | None:
    """The first type in the order integer, number, boolean, date, datetime that reaches 95 %."""
    for kind in ("integer", "number", "boolean", "date", "datetime"):
        for candidate in candidates:
            if candidate.type == kind and candidate.share >= GUESS_SHARE:
                return candidate
    return None


def analyse_column(position: int, raw: str, normalised: str, values: Sequence[object]) -> ColumnAnalysis:
    result = ColumnAnalysis(position + 1, raw, normalised)
    present: list[object] = []
    counter: Counter[str] = Counter()
    distinct: set[str] = set()
    null_like: Counter[str] = Counter()
    for value in values:
        if is_blank(value):
            continue
        text = canonical_str(value)
        result.non_empty += 1
        if isinstance(value, str):
            result.whitespace += has_whitespace_issue(value)
            result.non_nfc += not is_nfc(value)
            result.control += has_control(value)
            cleaned = clean_text(value, DEFAULT_TEXT)
            if cleaned.casefold() in NULL_LIKE:
                null_like[cleaned] += 1
                continue
            present.append(cleaned)
        else:
            present.append(value)
        counter[text] += 1
        if not result.distinct_capped:
            distinct.add(text)
            if len(distinct) > DISTINCT_CAP:
                result.distinct_capped = True
                distinct = set()
        length = len(text)
        result.min_length = length if result.min_length is None else min(result.min_length, length)
        result.max_length = length if result.max_length is None else max(result.max_length, length)
    result.distinct = DISTINCT_CAP if result.distinct_capped else len(distinct)
    result.null_like = dict(sorted(null_like.items()))
    result.top = [
        [value, count] for value, count in sorted(counter.items(), key=lambda i: (-i[1], i[0]))[:TOP]
    ]
    result.examples = list(dict.fromkeys(canonical_str(value) for value in present))[:EXAMPLES]
    if not present:
        return result
    distinct_values = weighted(present)
    total = len(present)
    numbers, number_notes, result.styles = _number_candidates(distinct_values, total)
    dates, date_notes = _date_candidate(distinct_values, total, BUILTIN_DATE_PATTERNS, "date")
    datetimes, datetime_notes = _date_candidate(distinct_values, total, BUILTIN_DATETIME_PATTERNS, "datetime")
    result.candidates = [*numbers, _boolean_candidate(distinct_values, total), dates, datetimes]
    result.guess = choose_guess(result.candidates)
    if result.guess is not None and result.guess.type in ("integer", "number"):
        result.notes.extend(number_notes)
    if result.guess is not None and result.guess.type == "date":
        result.notes.extend(date_notes)
    if result.guess is not None and result.guess.type == "datetime":
        result.notes.extend(datetime_notes)
    return result


def workbook_sheets(path: Path) -> list[tuple[str, int | None, int]]:
    """Every sheet with its detected header row and the number of non-empty rows below it."""
    try:
        book: Any = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:  # openpyxl raises many types for broken files
        raise InputError(f"{path.name}: not a readable workbook ({exc})") from exc
    try:
        sheets: list[tuple[str, int | None, int]] = []
        for name in book.sheetnames:
            worksheet = book[name]
            worksheet.reset_dimensions()
            rows = [list(row) for row in worksheet.iter_rows(values_only=True)]
            header = detect_header_row(rows)
            below = rows[header:] if header is not None else []
            records = sum(1 for row in below if any(not is_blank(value) for value in row))
            sheets.append((str(name), header, records))
        return sheets
    finally:
        book.close()


def select_sheets(sheets: Sequence[tuple[str, int | None, int]]) -> list[str]:
    """The sheets sharing the most common detected header row (ties: the earlier row)."""
    rows = Counter(header for _, header, _ in sheets if header is not None)
    if not rows:
        return [sheets[0][0]] if sheets else []
    common = max(rows.items(), key=lambda item: (item[1], -item[0]))[0]
    return [name for name, header, _ in sheets if header == common]


def analyse_file(
    path: Path,
    sheets: Sequence[str] = (),
    encoding: str = "auto",
    delimiter: str = "auto",
    header_row: int | str = "auto",
    max_rows: int = DEFAULT_MAX_ROWS,
) -> Analysis:
    all_sheets: list[tuple[str, int | None, int]] = []
    selected: list[str] = []
    patterns: tuple[str, ...] | None = None
    if path.suffix.lower() in (".xlsx", ".xlsm") and path.is_file():
        all_sheets = workbook_sheets(path)
        selected = list(sheets) if sheets else select_sheets(all_sheets)
        patterns = tuple(selected) if sheets else tuple(glob.escape(name) for name in selected)
    options = ReadOptions(encoding=encoding, delimiter=delimiter, header_row=header_row, sheets=patterns)
    source = open_source(path, options)
    info = source.info
    if info.kind == "xlsx":
        selected = [sheet.name for sheet in info.sheets]
    chunks = source.chunks(max_rows + 1 if info.kind == "csv" else None)
    chunk = next(iter(chunks))
    frame = chunk.frame
    stopped = row_count(frame) > max_rows
    if stopped:
        frame = frame.iloc[:max_rows]
    width = len(source.headers)
    names = data_columns(frame)
    rows = rows_of(frame, names)
    empty = sum(1 for row in rows if all(is_blank(value) for value in row))
    digests: set[bytes] = set()
    duplicates = 0
    for row in rows:
        if all(is_blank(value) for value in row):
            continue
        digest = hashlib.blake2b(repr([canonical_str(v) for v in row]).encode(), digest_size=16).digest()
        if digest in digests:
            duplicates += 1
        digests.add(digest)
    normalised = normalise_headers(source.headers)
    bases = header_bases(source.headers)
    repeated = sorted({raw for raw, base in zip(source.headers, bases, strict=True) if bases.count(base) > 1})
    columns = [
        analyse_column(
            position, source.headers[position], normalised[position], column(frame, names[position])
        )
        for position in range(width)
    ]
    empty_columns = [item.normalised for item in columns if item.non_empty == 0]
    votes: Counter[str] = Counter()
    styles: dict[str, tuple[str, str | None]] = {}
    for item in columns:
        guess = item.guess
        if guess is None or guess.type not in ("integer", "number") or guess.decimal is None:
            continue
        # styles maps each style to (number share, integer share); a column every style parses equally
        # well says nothing about the file's locale and does not vote
        which = 0 if guess.type == "number" else 1
        shares = {style: pair[which] for style, pair in item.styles.items()}
        if (
            len(set(shares.values())) <= 1
            and guess.share == 1.0
            and (guess.decimal, guess.thousands) == (".", None)
        ):
            continue
        key = _style_key(guess.decimal, guess.thousands)
        votes[key] += 1
        styles[key] = (guess.decimal, guess.thousands)
    locale = styles[max(votes.items(), key=lambda kv: kv[1])[0]] if votes else (".", None)
    for item in columns:
        prefer_locale(item, locale)
    date_formats = _global_patterns(columns, "date")
    datetime_formats = _global_patterns(columns, "datetime")
    warnings: list[str] = []
    if info.encoding_method in ("detected", "western"):
        warnings.append(
            f"encoding {info.encoding} was detected (confidence {info.encoding_confidence}); "
            "pin it with input.encoding"
        )
    if info.kind == "csv" and info.delimiter is None:
        warnings.append("no delimiter found; the file reads as a single column")
    if stopped:
        warnings.append(f"stopped after {max_rows} records (--max-rows)")
    return Analysis(
        info=info,
        headers=list(source.headers),
        columns=columns,
        analysed=row_count(frame),
        empty=empty,
        stopped_early=stopped,
        duplicate_records=duplicates,
        duplicate_headers=repeated,
        empty_columns=empty_columns,
        locale=locale,
        locale_votes=dict(votes),
        date_formats=date_formats,
        datetime_formats=datetime_formats,
        all_sheets=all_sheets,
        selected_sheets=selected,
        warnings=warnings,
    )


def prefer_locale(item: ColumnAnalysis, locale: tuple[str, str | None]) -> None:
    """Read a numeric column with the file's locale when that parses it just as well as its own best style."""
    guess = item.guess
    if guess is None or guess.type not in ("integer", "number") or locale not in item.styles:
        return
    number_share, integer_share = item.styles[locale]
    if (integer_share if guess.type == "integer" else number_share) >= guess.share:
        guess.decimal, guess.thousands = locale


def _global_patterns(columns: Sequence[ColumnAnalysis], kind: str) -> list[str]:
    """Patterns used by the guessed date or datetime columns, the most used column's first."""
    ordered: list[str] = []
    for item in sorted(columns, key=lambda c: -(c.non_empty)):
        if item.guess is not None and item.guess.type == kind:
            for pattern in item.guess.patterns:
                if pattern not in ordered and pattern != "native":
                    ordered.append(pattern)
    return ordered


def analysis_report(analysis: Analysis) -> dict[str, Any]:
    from cleandata.report.build import input_section, tool_info

    info = analysis.info
    decimal, thousands = analysis.locale
    return {
        "schema_version": 1,
        "kind": "profile",
        "title": f"Profile of {info.name}",
        "tool": tool_info(),
        "input": input_section(info),
        "records": {
            "analysed": analysis.analysed,
            "empty": analysis.empty,
            "short": info.short_records,
            "malformed": info.malformed,
            "duplicates": analysis.duplicate_records,
            "stopped_early": analysis.stopped_early,
        },
        "sheets": [
            {
                "name": name,
                "header_row": header,
                "records": records,
                "selected": name in analysis.selected_sheets,
            }
            for name, header, records in analysis.all_sheets
        ],
        "columns": [item.as_dict() for item in analysis.columns],
        "file_notes": {
            "duplicate_headers": analysis.duplicate_headers,
            "empty_columns": analysis.empty_columns,
            "duplicate_records": analysis.duplicate_records,
            "locale": {"decimal": decimal, "thousands": thousands, "votes": analysis.locale_votes},
            "date_formats": analysis.date_formats,
            "datetime_formats": analysis.datetime_formats,
        },
        "warnings": [{"code": "profile.note", "message": text, "count": 1} for text in analysis.warnings],
        "banner": None,
        "footer": None,
    }


TABLE_WIDTH: Final = 26


def _fit(text: str) -> str:
    """One table cell: line breaks, tabs and other invisible characters escaped, cut to the column width."""
    shown = display(text)
    return shown if len(shown) <= TABLE_WIDTH else shown[: TABLE_WIDTH - 1] + "~"


def text_table(analysis: Analysis) -> str:
    """The table ``clean-data profile`` prints to stdout."""
    info = analysis.info
    lines = [f"file       {display(info.name)} · {info.size:,} bytes · sha256 {info.sha256[:16]}"]
    if info.kind == "csv":
        confidence = "" if info.encoding_confidence is None else f", confidence {info.encoding_confidence}"
        lines.append(f"encoding   {info.encoding} ({info.encoding_method}{confidence})")
        delimiter = "none" if info.delimiter is None else repr(info.delimiter)
        share = "" if info.delimiter_share is None else f" (consistency {info.delimiter_share:.2f})"
        lines.append(f"delimiter  {delimiter}{share}")
        lines.append(f"header     row {info.header_row}")
    else:
        for name, header, records in analysis.all_sheets:
            mark = "*" if name in analysis.selected_sheets else " "
            row = "none" if header is None else f"row {header}"
            lines.append(f"sheet    {mark} {display(name)}: header {row}, {records} records")
    lines.append(
        f"records    {analysis.analysed:,} · empty {analysis.empty} · short {info.short_records} "
        f"· malformed {info.malformed} · duplicate {analysis.duplicate_records}"
    )
    decimal, thousands = analysis.locale
    lines.append(f"locale     {_style_key(decimal, thousands)}")
    lines.append("")
    heading = f"{'#':>3}  {'header':<26} {'normalised':<26} {'non-empty':>9} {'distinct':>8}  guess"
    lines.append(heading)
    for item in analysis.columns:
        guess = "string" if item.guess is None else f"{item.guess.label()} {item.guess.share * 100:.1f} %"
        if item.non_empty == 0:
            guess = "empty"
        distinct = f"{item.distinct}+" if item.distinct_capped else str(item.distinct)
        lines.append(
            f"{item.position:>3}  {_fit(item.raw):<26} {_fit(item.normalised):<26} "
            f"{item.non_empty:>9} {distinct:>8}  {guess}"
        )
        for token, count in item.null_like.items():
            lines.append(f"{'':>5}null-like {token!r} x{count}")
        for note in item.notes:
            lines.append(f"{'':>5}note: {note}")
    if analysis.duplicate_headers:
        lines.append("")
        lines.append("repeated headers: " + ", ".join(display(name) for name in analysis.duplicate_headers))
    if analysis.empty_columns:
        lines.append("empty columns: " + ", ".join(analysis.empty_columns))
    for warning in analysis.warnings:
        lines.append(f"warning: {warning}")
    return "\n".join(lines) + "\n"
